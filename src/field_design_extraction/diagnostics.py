"""Help for setting up the block ranges: what the map holds and where."""

import numpy as np
import pandas as pd
from openpyxl.utils import get_column_letter


def color_summary(cells):
    """One row per fill color: how many cells, what they hold and where they are."""
    rows = []

    for fill, group in cells.groupby("fill", dropna=False):
        with_value = group[group["value"].notna()]

        rows.append({
            "fill": fill,
            "n_cells": len(group),
            "n_with_value": len(with_value),
            "n_numeric": int(group["value_type"].isin(["numeric", "numeric_text"]).sum()),
            "n_text": int(group["value_type"].isin(["text", "alphanumeric"]).sum()),
            "n_merged_anchors": int((group["is_merged"] & group["is_anchor"]).sum()),
            "rows": f"{group['row'].min()}-{group['row'].max()}",
            "columns": (
                f"{get_column_letter(group['column'].min())}-"
                f"{get_column_letter(group['column'].max())}"
            ),
            "examples": with_value["value"].head(5).tolist(),
        })

    return (
        pd.DataFrame(rows)
        .sort_values("n_cells", ascending=False)
        .reset_index(drop=True)
    )


def plot_cells(cells, design=None, title=None):
    """Map of the cells colored by fill; with `design`, plot cells are marked by block.

    Use it to check that every block range covers its plots and nothing else.
    """
    import matplotlib.pyplot as plt

    min_row, max_row = cells["row"].min(), cells["row"].max()
    min_col, max_col = cells["column"].min(), cells["column"].max()

    fills = [f for f in cells["fill"].dropna().unique()]
    code = {f: i + 1 for i, f in enumerate(fills)}

    matrix = np.zeros((max_row - min_row + 1, max_col - min_col + 1))

    for row, column, fill in zip(cells["row"], cells["column"], cells["fill"]):
        matrix[row - min_row, column - min_col] = code.get(fill, 0)

    width = min(18, max(8, matrix.shape[1] * 0.18))
    height = min(14, max(5, matrix.shape[0] * 0.18))
    fig, ax = plt.subplots(figsize=(width, height))

    ax.imshow(
        matrix,
        cmap=plt.get_cmap("Pastel1", max(len(fills) + 1, 2)),
        interpolation="nearest",
        aspect="auto",
        extent=(min_col - 0.5, max_col + 0.5, max_row + 0.5, min_row - 0.5),
    )

    if design is not None:
        markers = plt.get_cmap("tab10")

        for i, (block_id, block) in enumerate(design.groupby("block_id", sort=False)):
            ax.scatter(
                block["column"], block["row"],
                s=12, color=markers(i % 10), label=f"{block_id} ({len(block)})",
            )

        ax.legend(loc="upper left", bbox_to_anchor=(1.01, 1), fontsize=8)

    step = max(1, (max_col - min_col + 1) // 25)
    ticks = list(range(min_col, max_col + 1, step))
    ax.set_xticks(ticks)
    ax.set_xticklabels([get_column_letter(c) for c in ticks], rotation=90)
    ax.set_xlabel("Excel column")
    ax.set_ylabel("Excel row")
    ax.set_title(title or "Map cells by fill color")
    fig.tight_layout()

    return fig
