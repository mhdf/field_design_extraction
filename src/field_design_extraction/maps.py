"""Ratings marked on the field map itself: one sheet per date, status of every plot cell."""

import re

import pandas as pd
from openpyxl.utils import get_column_letter

from .workbook import read_cells


def _cell_text(value, ignore_text):
    """Cell value as stripped text; None when empty or matching ignore_text."""
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return None

    if isinstance(value, float) and value.is_integer():
        value = int(value)

    text = str(value).strip()

    if not text or (ignore_text and re.fullmatch(ignore_text, text)):
        return None

    return text


def _rule_matches(rule, text, fill):
    """All conditions present in the rule hold (text: regex, case-insensitive)."""
    if "empty" in rule and (text is None) != rule["empty"]:
        return False

    if "text" in rule and (text is None or not re.search(rule["text"], text, re.IGNORECASE)):
        return False

    if "fills" in rule and fill not in rule["fills"]:
        return False

    return True


def classify_cell(text, fill, rules, default_status="healthy"):
    """Status of one cell: the first matching rule; default_status when empty, 'unknown' otherwise."""
    for rule in rules:
        if _rule_matches(rule, text, fill):
            return rule["status"]

    return default_status if text is None else "unknown"


def read_map_ratings(source, design):
    """One row per design plot and date, with its status read from that date's sheet.

    source: dict with workbook, disease, assessments ({name: {sheet, date}}),
            status_rules and optional ignore_text, default_status, flags
            (see the README).
    """
    rules = source["status_rules"]
    ignore_text = source.get("ignore_text")
    default_status = source.get("default_status", "healthy")
    flags = source.get("flags", {})

    region = (
        f"{get_column_letter(design['column'].min())}{design['row'].min()}:"
        f"{get_column_letter(design['column'].max())}{design['row'].max()}"
    )

    tables = []

    for key, assessment in source["assessments"].items():
        cells = read_cells(source["workbook"], assessment["sheet"], region)

        table = design[["plot_key", "block_id", "plot_number", "row", "column"]].merge(
            cells[["row", "column", "value", "fill"]],
            on=["row", "column"],
            how="left",
        )

        texts = [_cell_text(value, ignore_text) for value in table["value"]]

        table["status"] = [
            classify_cell(text, fill, rules, default_status)
            for text, fill in zip(texts, table["fill"])
        ]

        # --- flags: True when the cell fill is one of the listed colors
        for name, fills in flags.items():
            table[name] = table["fill"].isin(fills)

        table["date"] = pd.Timestamp(assessment["date"]).date().isoformat()
        table["disease"] = source["disease"]
        table["assessment"] = key
        table["link_status"] = "linked"

        tables.append(table.drop(columns=["row", "column"]))

    return pd.concat(tables, ignore_index=True)


def map_value_columns(source):
    """Columns written after RATING_COLUMNS for a map source."""
    return ["status", *source.get("flags", {}), "value", "fill"]


def report_map_ratings(source_id, ratings, source):
    """Print plots per status for every date and list codes no rule covers."""
    print()
    print("=" * 80)
    print(f"MAP RATINGS — {source_id}")
    print("=" * 80)

    table = (
        ratings.pivot_table(index="date", columns="status", values="plot_key", aggfunc="count", fill_value=0)
    )

    for name in source.get("flags", {}):
        table[name] = ratings.groupby("date")[name].sum()

    print(table.to_string())

    unknown = ratings[ratings["status"] == "unknown"]

    if not unknown.empty:
        print()
        print("WARNING: codes not covered by status_rules (status 'unknown'):")
        print(
            unknown.groupby(["value", "fill"], dropna=False)
            .size()
            .rename("n_cells")
            .reset_index()
            .to_string(index=False)
        )
