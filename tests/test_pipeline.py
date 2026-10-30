"""Tests on a small made-up workbook: map sheet + two kinds of ratings sheets."""

import contextlib
import io
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

import field_design_extraction as fde


def make_workbook(path):
    wb = Workbook()

    # --- map: trial A (B3:D4, plots 1-6), trial B (E3:F4, plot 3 twice)
    ws = wb.active
    ws.title = "map"
    ws["B2"] = "Trial A"
    ws.merge_cells("B2:D2")
    ws["E2"] = "Trial B"
    ws.merge_cells("E2:F2")
    for column, values in zip("BCD", [(1, 4), (2, 5), (3, 6)]):
        ws[f"{column}3"], ws[f"{column}4"] = values
    ws["E3"], ws["F3"], ws["E4"], ws["F4"] = 1, 2, 3, 3
    ws["G3"] = "=1+1"   # formula with no saved value (file not saved by Excel)

    # --- one sheet per trial, header on row 3, MIN row below the table
    ws = wb.create_sheet("A")
    ws["A1"], ws["B1"] = "Data:", "2025-09-10"
    ws.append([])
    ws["A3"], ws["B3"], ws["C3"] = "Parcela", "Nota", "Obs"
    ws.append([1, 10, "x"])
    ws.append([2, 20, None])
    ws.append([9, 30, None])
    ws.append([None, 10, "MIN"])

    # --- one sheet per date, with a block column
    ws = wb.create_sheet("20250901")
    ws.append(["plot", "block", "score"])
    ws.append([1, "Trial B", 5])
    ws.append([3, "Trial B", 7])
    ws.append([1, "Unknown", 1])

    wb.save(path)


class PipelineTest(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.folder = Path(self.tmp.name)
        self.workbook = self.folder / "field.xlsx"
        make_workbook(self.workbook)

        self.layout = {
            "workbook": self.workbook,
            "sheet": "map",
            "region": "B2:G4",
            "output_folder": self.folder / "consolidated",
            "unit_rules": {"value_types": ["numeric"], "merged": False},
            "blocks": [
                {"block_id": "A", "label": "Trial A", "regions": ["B3:D4"]},
                {"block_id": "B", "label": "Trial B", "regions": ["E3:F4"]},
            ],
        }
        self.per_sheet = {
            "layout_id": "2025_design",
            "disease": "net_blotch",
            "workbook": self.workbook,
            "header_row": 3,
            "plot_column": "Parcela",
            "block_column": None,
            "value_columns": {"Nota": "score"},
            "assessments": {"A": {"sheet": "A", "date": "2025-09-10", "block": "A"}},
        }
        self.per_date = {
            "layout_id": "2025_design",
            "disease": "tan_spot",
            "workbook": self.workbook,
            "plot_column": "plot",
            "block_column": "block",
            "value_columns": {"score": "score"},
            "block_mapping": {"Trial B": "B"},
            "assessments": {"d1": {"sheet": "20250901", "date": "2025-09-01"}},
        }

    def tearDown(self):
        self.tmp.cleanup()

    def quiet(self, function, *args):
        with contextlib.redirect_stdout(io.StringIO()) as out:
            result = function(*args)
        return result, out.getvalue()

    def test_design_one_row_per_plot_cell(self):
        design = fde.build_design("2025_design", self.layout)

        self.assertEqual(len(design), 10)
        self.assertEqual(design["block_id"].value_counts().to_dict(), {"A": 6, "B": 4})
        self.assertEqual(design.loc[0, "plot_key"], "A:1:3:2")
        self.assertEqual(design.loc[0, "block_label"], "Trial A")

    def test_formula_cell_is_not_read_as_text(self):
        cells, out = self.quiet(fde.read_cells, self.workbook, "map", "G3")

        self.assertIsNone(cells.loc[0, "value"])
        self.assertEqual(cells.loc[0, "formula"], "=1+1")
        self.assertIn("no saved value", out)

    def test_overlapping_blocks_raise(self):
        self.layout["blocks"][1]["regions"] = ["D3:F4"]

        with self.assertRaisesRegex(ValueError, "more than one block"):
            fde.build_design("2025_design", self.layout)

    def test_ratings_one_sheet_per_block(self):
        ratings, out = self.quiet(fde.read_ratings, self.per_sheet)

        self.assertIn("dropped 1 rows", out)
        self.assertEqual(ratings["plot_number"].tolist(), ["1", "2", "9"])
        self.assertEqual(ratings["block_id"].unique().tolist(), ["A"])
        self.assertEqual(ratings["score"].tolist(), [10, 20, 30])
        self.assertEqual(ratings["date"].unique().tolist(), ["2025-09-10"])

    def test_link_status(self):
        design = fde.build_design("2025_design", self.layout)
        per_sheet = fde.link_ratings(self.quiet(fde.read_ratings, self.per_sheet)[0], design)
        per_date = fde.link_ratings(fde.read_ratings(self.per_date), design)

        self.assertEqual(per_sheet["link_status"].tolist(), ["linked", "linked", "no_match"])
        self.assertEqual(per_sheet.loc[0, "plot_key"], "A:1:3:2")
        self.assertEqual(per_date["link_status"].tolist(), ["linked", "ambiguous", "unmapped_block"])

    def test_export_writes_design_and_one_file_per_date(self):
        layouts = {"2025_design": self.layout}
        sources = {"2025_net_blotch": self.per_sheet, "2025_tan_spot": self.per_date}

        self.quiet(fde.run, layouts, sources)

        out = self.folder / "consolidated"
        self.assertEqual(
            sorted(p.name for p in out.iterdir()),
            [
                "2025_design.csv",
                "2025_net_blotch_20250910.csv",
                "2025_net_blotch_20250910_unlinked.csv",
                "2025_tan_spot_20250901.csv",
                "2025_tan_spot_20250901_unlinked.csv",
            ],
        )

        rated = pd.read_csv(out / "2025_net_blotch_20250910.csv")
        self.assertEqual(
            list(rated.columns),
            ["plot_key", "block_id", "plot_number", "date", "disease", "score"],
        )
        self.assertEqual(len(rated), 2)

        # --- once every rating links, the old _unlinked file is removed
        self.per_sheet["assessments"]["A"]["sheet"] = "A"
        sources = {"2025_net_blotch": {**self.per_sheet, "block_mapping": {"A": "A"}}}
        design = fde.build_design("2025_design", self.layout)
        linked = fde.link_ratings(self.quiet(fde.read_ratings, sources["2025_net_blotch"])[0], design)
        linked = linked[linked["link_status"] == "linked"]
        fde.write_ratings(linked, out, "2025_net_blotch", ["score"])

        self.assertFalse((out / "2025_net_blotch_20250910_unlinked.csv").exists())

    def test_coverage_figure_for_ratings_table(self):
        import matplotlib
        matplotlib.use("Agg")

        layouts = {"2025_design": {**self.layout, "figure": {"label_rotation": 0, "gap_rule": "disjoint"}}}
        sources = {"2025_net_blotch": self.per_sheet}

        with contextlib.redirect_stdout(io.StringIO()):
            designs = fde.build_designs(layouts)
            ratings = fde.build_ratings(sources, designs)
            fde.save_figures(layouts, sources, designs, ratings, which=["coverage"])

        self.assertEqual(
            [p.name for p in (self.folder / "consolidated" / "figures").iterdir()],
            ["2025_net_blotch_20250910_coverage.png"],
        )


class MapRatingsTest(unittest.TestCase):
    """Orchard-style map: every cell is a tree, one sheet per date, status from value and fill."""

    YELLOW, BLACK, GREEN = "FFFFFF00", "FF000000", "FF00FF00"

    def setUp(self):
        from openpyxl.styles import PatternFill

        self.tmp = tempfile.TemporaryDirectory()
        self.workbook = Path(self.tmp.name) / "orchard.xlsx"

        def paint(ws, coordinate, color):
            ws[coordinate].fill = PatternFill("solid", fgColor=color)

        wb = Workbook()
        for i, name in enumerate(["May 2019", "May 2020"]):
            ws = wb.active if i == 0 else wb.create_sheet()
            ws.title = name
            for row in range(1, 3):
                for column in "ABC":
                    paint(ws, f"{column}{row}", self.GREEN)
            ws["A1"] = "A1"        # plot label written in a tree cell
            if i == 1:
                ws["B1"] = 1       # new case
                paint(ws, "B1", self.YELLOW)
                paint(ws, "C1", self.BLACK)   # earlier case, no value
                ws["A2"] = "R2"
                ws["B2"] = "Gomose"
                ws["C2"] = "???"
        wb.save(self.workbook)

        self.layout = {
            "workbook": self.workbook,
            "sheet": "May 2019",
            "region": "A1:C2",
            "unit_rules": {"value_types": None, "include_empty": True, "unit_id": "coordinate"},
            "blocks": [{"block_id": "A1", "regions": ["A1:C2"], "attributes": {"treatment": "A"}}],
        }
        self.source = {
            "kind": "map",
            "layout_id": "orchard",
            "disease": "hlb",
            "workbook": self.workbook,
            "ignore_text": r"[A-F][1-6]",
            "status_rules": [
                {"text": r"gomo", "status": "gummosis"},
                {"text": r"^r\s*\d*$", "status": "hlb"},
                {"text": r"^1$", "status": "hlb"},
                {"empty": True, "fills": [self.BLACK, self.YELLOW], "status": "hlb"},
            ],
            "flags": {"new_case": [self.YELLOW]},
            "assessments": {
                "2019-05": {"sheet": "May 2019", "date": "2019-05-01"},
                "2020-05": {"sheet": "May 2020", "date": "2020-05-01"},
            },
        }

    def tearDown(self):
        self.tmp.cleanup()

    def test_every_tree_is_a_plot_with_block_attributes(self):
        design = fde.build_design("orchard", self.layout)

        self.assertEqual(len(design), 6)
        self.assertEqual(design.loc[0, "plot_key"], "A1:A1:1:1")
        self.assertEqual(design["treatment"].unique().tolist(), ["A"])

    def test_status_per_tree_and_date(self):
        design = fde.build_design("orchard", self.layout)
        with contextlib.redirect_stdout(io.StringIO()) as out:
            ratings = fde.read_map_ratings(self.source, design)
            fde.report_map_ratings("orchard_hlb", ratings, self.source)

        first = ratings[ratings["date"] == "2019-05-01"]
        second = ratings[ratings["date"] == "2020-05-01"].set_index("plot_number")

        self.assertEqual(len(ratings), 12)
        self.assertEqual(first["status"].unique().tolist(), ["healthy"])
        self.assertEqual(
            second["status"].to_dict(),
            {"A1": "healthy", "B1": "hlb", "C1": "hlb", "A2": "hlb", "B2": "gummosis", "C2": "unknown"},
        )
        self.assertEqual(second["new_case"].sum(), 1)
        self.assertTrue(second.loc["B1", "new_case"])
        self.assertIn("not covered by status_rules", out.getvalue())

    def test_figures_are_saved(self):
        import matplotlib
        matplotlib.use("Agg")

        layouts = {"orchard": {**self.layout, "output_folder": Path(self.tmp.name) / "consolidated",
                               "figure": {"color_by": "treatment", "plot_size": (1, 1), "block_outlines": True}}}
        self.source["figure"] = {"incidence": {"status": "hlb", "group": "treatment"}}
        sources = {"hlb": self.source}

        with contextlib.redirect_stdout(io.StringIO()):
            designs = fde.build_designs(layouts)
            ratings = fde.build_ratings(sources, designs)
            fde.save_figures(layouts, sources, designs, ratings, which=fde.FIGURE_TYPES)

        self.assertEqual(
            sorted(p.name for p in (Path(self.tmp.name) / "consolidated" / "figures").iterdir()),
            ["hlb_20190501_map.png", "hlb_20200501_map.png", "hlb_distribution.png",
             "hlb_incidence.png", "orchard.png"],
        )

    def test_dates_and_date_labels(self):
        from field_design_extraction.ratings import date_label, date_stem

        self.assertEqual(date_label("2025-9-1"), "2025-09-01")
        self.assertEqual(date_label("01/09/2025"), "2025-09-01")
        self.assertEqual(date_label(pd.Timestamp("2025-09-01 10:00")), "2025-09-01")
        self.assertEqual(date_label(" Date 1 "), "Date 1")
        self.assertEqual(date_stem("2025-09-01"), "20250901")
        self.assertEqual(date_stem("Date 1"), "date_1")

    def test_other_classes_are_grouped_with_their_names(self):
        from field_design_extraction.figures import group_other_values

        values, label = group_other_values(
            pd.Series(["healthy", "hlb", "dead", "no_tree", None, "dead"]),
            keep=["healthy", "hlb"], order=["healthy", "hlb", "no_tree", "dead"],
        )

        self.assertEqual(label, "others (no tree, dead)")
        self.assertEqual(values.tolist()[:4], ["healthy", "hlb", label, label])
        # --- a missing value stays missing (None with pandas 2, NaN with pandas 3)
        self.assertTrue(pd.isna(values.iloc[4]))

    def test_many_numeric_values_become_classes(self):
        from field_design_extraction.figures import score_classes

        classes, labels = score_classes(pd.Series([12, 25, 25.5, 68, None] + list(range(30, 45))))

        self.assertEqual(labels, ["10–20", "20–30", "30–40", "40–50", "50–60", "60–70"])
        self.assertEqual(classes.iloc[:5].tolist()[:4], ["10–20", "20–30", "20–30", "60–70"])
        self.assertTrue(pd.isna(classes.iloc[4]))

        # --- the last class includes its upper edge
        classes, labels = score_classes(pd.Series([0, 55, 100]), bins=list(range(0, 101, 10)))
        self.assertEqual(classes.tolist(), ["0–10", "50–60", "90–100"])


if __name__ == "__main__":
    unittest.main()
