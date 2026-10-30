"""Create the two synthetic example workbooks in examples/data (fixed random seed).

Scores (0-100) come from a smooth random surface over the field map, so that
neighbouring plots get similar values, as disease often does in the field.
No real plot numbers, genotypes or treatments are used.

    python examples/create_example_workbooks.py
"""

from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DATA = Path(__file__).resolve().parent / "data"
RNG = np.random.default_rng(2026)


def smooth_field(shape, sigma):
    """Random surface (values 0-1) with spatial aggregation: Gaussian-smoothed noise."""
    radius = int(3 * sigma)
    kernel = np.exp(-0.5 * (np.arange(-radius, radius + 1) / sigma) ** 2)
    kernel /= kernel.sum()

    noise = RNG.normal(size=(shape[0] + 2 * radius, shape[1] + 2 * radius))
    noise = np.apply_along_axis(np.convolve, 0, noise, kernel, mode="same")
    noise = np.apply_along_axis(np.convolve, 1, noise, kernel, mode="same")
    field = noise[radius:-radius, radius:-radius]

    return (field - field.min()) / (field.max() - field.min())


def fill(color):
    return PatternFill("solid", fgColor=color)


def serpentine(n_rows, n_columns):
    """Plot order from the bottom-left corner, alternating direction in each row."""
    order = []
    for i, row in enumerate(range(n_rows - 1, -1, -1)):
        columns = range(n_columns) if i % 2 == 0 else range(n_columns - 1, -1, -1)
        order += [(row, column) for column in columns]
    return order


# =============================================
# EXAMPLE 1 — plots in merged cells, ratings table with a block column
# =============================================

def example_1():
    """Map: 4 blocks of plots, each plot 2 merged cells tall; ratings: one sheet per date."""
    # --- block: (label, first map row, first map column, plot rows, plot columns, first plot number, color)
    blocks = [
        ("Block A", 4, 2, 10, 6, 101, "FFF2B134"),
        ("Block B", 4, 9, 10, 6, 201, "FFF28E2B"),
        ("Block C", 4, 16, 4, 10, 301, "FF43B97F"),
        ("Block D", 14, 16, 5, 10, 401, "FF36A9D6"),
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "field_map"
    ws["B1"] = "Example field — plots of 2 x 5 m, numbered in serpentine order"
    ws["B1"].font = Font(bold=True)

    field = smooth_field((30, 30), sigma=3.0)
    plots = []   # (block label, plot number, map row, map column)

    for label, row0, col0, n_rows, n_cols, first, color in blocks:
        ws.cell(row0 - 1, col0, label).font = Font(bold=True)
        ws.merge_cells(start_row=row0 - 1, start_column=col0, end_row=row0 - 1, end_column=col0 + n_cols - 1)

        for k, (r, c) in enumerate(serpentine(n_rows, n_cols)):
            row, column = row0 + 2 * r, col0 + c
            ws.merge_cells(start_row=row, start_column=column, end_row=row + 1, end_column=column)
            cell = ws.cell(row, column, first + k)
            cell.fill = fill(color)
            cell.alignment = Alignment(horizontal="center", vertical="center", text_rotation=90)
            plots.append((label, first + k, row, column))

    for column in range(1, 27):
        ws.column_dimensions[get_column_letter(column)].width = 5

    # --- ratings: one sheet per assessment date, block labels in capitals; the assessment
    #     grows over time: date 2 repeats every plot of date 1 and adds new ones
    order = RNG.permutation(len(plots))

    for sheet_name, share, increase in [("date_1", 0.4, 0.0), ("date_2", 0.7, 25.0)]:
        sheet = wb.create_sheet(sheet_name)
        sheet.append(["plot", "block", "score", "evaluator"])
        chosen = order[: int(share * len(plots))]

        for i in sorted(chosen):
            label, number, row, column = plots[i]
            score = 75 * field[row - 3, column - 1] + increase + RNG.normal(0, 4)
            sheet.append([number, label.upper(), int(np.clip(round(score), 0, 100)), "Evaluator 1"])

    DATA.mkdir(exist_ok=True)
    wb.save(DATA / "example_1_merged_plots.xlsx")
    print(f"example 1: {len(plots)} plots")


# =============================================
# EXAMPLE 2 — one cell per plot, one ratings sheet per trial
# =============================================

def example_2():
    """Map: 4 trials side by side, plots numbered from 1 in each trial; ratings: one sheet per trial."""
    # --- trial: (name, plot columns, color); all trials have 8 plot rows
    trials = [
        ("Trial 1", 5, "FFC7D9EA"),
        ("Trial 2", 4, "FFB5D5C3"),
        ("Trial 3", 10, "FFE5D8A8"),
        ("Trial 4", 6, "FFE6B69E"),
    ]
    n_rows, first_row = 8, 4

    wb = Workbook()
    ws = wb.active
    ws.title = "field_map"
    ws["B1"] = "Example field — 4 trials, plots numbered from 1 within each trial"
    ws["B1"].font = Font(bold=True)

    total_columns = sum(t[1] for t in trials)
    field = smooth_field((n_rows, total_columns), sigma=1.6)
    trial_plots = {}
    column0 = 2

    for name, n_cols, color in trials:
        ws.cell(first_row - 1, column0, name).font = Font(bold=True)
        ws.merge_cells(start_row=first_row - 1, start_column=column0,
                       end_row=first_row - 1, end_column=column0 + n_cols - 1)

        plots = []
        for k, (r, c) in enumerate(serpentine(n_rows, n_cols)):
            cell = ws.cell(first_row + r, column0 + c, k + 1)
            cell.fill = fill(color)
            cell.alignment = Alignment(horizontal="center")
            plots.append((k + 1, r, column0 + c - 2))

        trial_plots[name] = plots
        column0 += n_cols

    # --- ratings: metadata rows, header on row 7, MIN/MAX rows below the table
    for name, plots in trial_plots.items():
        sheet = wb.create_sheet(name.replace(" ", "_"))
        sheet["A1"], sheet["B1"] = "Date:", "Date 1"
        sheet["A2"], sheet["B2"] = "Growth stage:", 71
        sheet["A3"], sheet["B3"] = "Assessed area:", "Plot"
        sheet["A7"], sheet["B7"] = "Plot", "Severity (%)"
        sheet.column_dimensions["A"].width = 15
        sheet.column_dimensions["B"].width = 13

        for number, r, c in sorted(plots):
            score = 10 + 85 * field[r, c] + RNG.normal(0, 3)
            sheet.append([number, int(np.clip(round(score), 0, 100))])

        last = sheet.max_row
        sheet.append([None, f"=MIN(B8:B{last})", "min"])
        sheet.append([None, f"=MAX(B8:B{last})", "max"])

    wb.save(DATA / "example_2_one_sheet_per_trial.xlsx")
    print(f"example 2: {sum(len(p) for p in trial_plots.values())} plots")


if __name__ == "__main__":
    example_1()
    example_2()
